from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from larp_studio.project import Project
from larp_studio.exports import export_xlsx, export_pdf


@pytest.fixture
def game(tmp_path):
    p = Project.create(tmp_path / 'Проект', 'Русская игра')
    loc = p.add_location('Библиотека')
    p.add_record('Тайная записка', loc, [{'kind': 'text', 'value': 'Секрет'}])
    yield p
    p.close()


def test_excel_contains_codes_locations_types_and_russian(game, tmp_path):
    output = tmp_path / 'Коды.xlsx'
    export_xlsx(game, output)
    workbook = load_workbook(output)
    sheet = workbook.active
    assert list(sheet.values)[0][:4] == ('Код', 'Название', 'Локация', 'Содержимое')
    row = list(sheet.values)[1]
    assert row[:4] == (game.records()[0]['code'], 'Тайная записка', 'Библиотека', 'Текст')


def test_excel_does_not_execute_user_text_as_formula(game, tmp_path):
    game.add_record('=HYPERLINK("bad")')
    output = tmp_path / 'Коды.xlsx'
    export_xlsx(game, output)
    sheet = load_workbook(output).active
    for row in sheet.iter_rows(min_row=2):
        assert row[1].data_type == 's'


def test_pdf_default_hides_spoilers_and_contains_code(game, tmp_path):
    output = tmp_path / 'Карточки.pdf'
    export_pdf(game, output)
    pdf = PdfReader(output)
    assert len(pdf.pages) == 1
    text = pdf.pages[0].extract_text()
    assert game.records()[0]['code'] in text
    assert 'Тайная записка' not in text
    assert 'Библиотека' not in text


def test_pdf_optional_russian_labels_and_a4_pagination(game, tmp_path):
    for i in range(12):
        game.add_record(f'Карточка {i}')
    output = tmp_path / 'Карточки.pdf'
    export_pdf(game, output, show_title=True, show_location=True)
    pdf = PdfReader(output)
    assert len(pdf.pages) == 2
    text = '\n'.join(p.extract_text() for p in pdf.pages)
    assert 'Тайная записка' in text
    assert 'Библиотека' in text
    assert abs(float(pdf.pages[0].mediabox.width) - 595.28) < 1
