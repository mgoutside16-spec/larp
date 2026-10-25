"""Print exports; no dependency on a running QApplication."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from .files import atomic_output

KINDS = {'text': 'Текст', 'image': 'Фото', 'audio': 'Аудио', 'video': 'Видео'}


def material_types(row):
    return ', '.join(KINDS[k] for k in KINDS if any(b['kind'] == k for b in row['blocks']))


def ordered_records(project):
    return sorted(project.records(), key=lambda r: ((r['location'] or '').casefold(), r['title'].casefold(), r['code']))


def export_xlsx(project, destination):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Коды'
    sheet.append(['Код', 'Название', 'Локация', 'Содержимое', 'Изменено'])
    for row in ordered_records(project):
        sheet.append([row['code'], row['title'], row['location'] or '', material_types(row), row['modified']])
        # User text must never become an Excel formula.
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='315EA8')
    for column, width in [('A', 14), ('B', 45), ('C', 30), ('D', 30), ('E', 28)]:
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    with atomic_output(destination) as temp:
        workbook.save(temp)
    workbook.close()


def _font():
    name = 'LARPDejaVu'
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(Path(__file__).parent / 'assets' / 'DejaVuSans.ttf')))
    return name


def _fit_label(canvas, value, x, y, width, font):
    value = value.replace('\n', ' ')
    while pdfmetrics.stringWidth(value, font, 10) > width and len(value) > 1:
        value = value[:-2] + '…'
    canvas.setFont(font, 10)
    canvas.drawCentredString(x, y, value)


def export_pdf(project, destination, *, show_title=False, show_location=False):
    font = _font()
    width, height = A4
    margin, gap = 32, 12
    card_width = (width - 2 * margin - gap) / 2
    card_height = (height - 2 * margin - 3 * gap) / 4
    with atomic_output(destination) as temp:
        canvas = Canvas(str(temp), pagesize=A4)
        canvas.setTitle('Коды для печати — ' + project.info['name'])
        rows = ordered_records(project)
        for i, row in enumerate(rows):
            if i and i % 8 == 0:
                canvas.showPage()
            index = i % 8
            x = margin + (index % 2) * (card_width + gap)
            y = height - margin - (index // 2 + 1) * card_height - (index // 2) * gap
            canvas.setStrokeColorRGB(.78, .82, .88)
            canvas.rect(x, y, card_width, card_height)
            center = x + card_width / 2
            canvas.setFillColorRGB(.12, .18, .27)
            canvas.setFont(font, 43)
            canvas.drawCentredString(center, y + card_height / 2, row['code'])
            if show_title:
                _fit_label(canvas, row['title'], center, y + 43, card_width - 20, font)
            if show_location:
                _fit_label(canvas, row['location'] or 'Без локации', center, y + 25, card_width - 20, font)
        if not rows:
            canvas.setFont(font, 14)
            canvas.drawString(margin, height - margin - 20, 'В проекте пока нет записей.')
        canvas.save()
